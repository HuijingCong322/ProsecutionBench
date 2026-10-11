"""Strategy review sheet v2 (no reviewer suggestions): per gold app, the automatic label plus the evidence needed
to check it yourself: every A1 document with USPTO's official document description (catalog.csv) and links to
the A1 claim listing, the A1 remarks and the claim listing in force before A1 (for cancellations).
The week-3 update plan asks for manual correction of the automatic labels. Output: gold/response_strategy_review_v2.xlsx"""
import csv, glob, json, os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
A = Path(__file__).resolve().parent.parent; R = A.parent; S = R / "data_v1_share"
MAC = "file:///Users/fionay/examinerbench/"
apps = sorted(r["app_id"] for r in csv.DictReader(open(S / "labels.csv")) if r["gold_subset"] == "True")
wb = Workbook(); ws = wb.active; ws.title = "strategy"
ws.append(["app_id", "A1 日期", "A1 提交的全部文件（USPTO 官方文件名称）", "A1 claims 清单", "A1 意见陈述 (REM)", "A1 修改书 (A...)",
           "A1 之前最后一份 claims 清单（查删除用）", "自动标签 (auto)", "你的判断 ok / fix", "fix 时写正确标签（分号分隔）", "依据（简短，例如 'CLM 第 3 条 Currently Amended'）"])
for c in ws[1]: c.font = Font(bold=True)
def link(cell, path, label):
    if path: cell.value = label; cell.hyperlink = MAC + str(Path(path).relative_to(R)); cell.font = Font(color="0563C1", underline="single")
for app in apps:
    e = json.loads((S / "examples" / f"{app}.json").read_text())
    cat = {r["doc_id"]: r for r in csv.DictReader(open(R / "data_v1" / "apps" / app / "catalog.csv"))}
    docs = "\n".join(f"{cat[d]['doc_code']} — {cat[d]['description']}" for d in e["a1"]["doc_ids"] if d in cat)
    def find(code):
        for d in e["a1"]["doc_ids"]:
            if cat.get(d, {}).get("doc_code") == code:
                f = [x for x in glob.glob(str(S / "apps" / app / "text" / f"*{d}*.txt")) if "substantive" not in x and not x.endswith(".json")]
                if f: return f[0]
        return None
    prev = sorted(x for x in glob.glob(str(S / "apps" / app / "text" / "*_CLM_*.txt")) if os.path.basename(x)[:10] < e["cutoff"])
    ws.append([app, e["cutoff"], docs, "", "", "", "", ";".join(sorted(e["labels"]["response_strategy"])), "", "", ""])
    r = ws.max_row
    link(ws.cell(r, 4), find("CLM"), "打开"); link(ws.cell(r, 5), find("REM"), "打开"); link(ws.cell(r, 6), find("A..."), "打开")
    link(ws.cell(r, 7), prev[-1] if prev else None, os.path.basename(prev[-1])[:14] if prev else "")
    if not find("CLM"): ws.cell(r, 4).value = "（A1 没有 claims 清单）"
for col, wd in zip("ABCDEFGHIJK", [10, 11, 60, 9, 9, 9, 16, 45, 10, 40, 40]): ws.column_dimensions[col].width = wd
for row in ws.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
g = wb.create_sheet("怎么判断")
for t in ["每个标签对应 A1 里能看到的一个具体事实（看'A1 提交的全部文件'这一列和链接的原文）：",
          "amend_claims：A1 claims 清单里至少一条标 (Currently Amended)。",
          "add_claims：A1 claims 清单里至少一条标 (New)。",
          "cancel_claims：A1 claims 清单里至少一条标 (Canceled)，而且在'A1 之前最后一份 claims 清单'里它还没被删除（之前已经删除、这次只是照抄 (Canceled) 的不算）。之前是否已删除，也可以看 O1 开头对 claims 状态的说明（例如 'Claims ... were cancelled'）。",
          "argue：有 REM（Applicant Arguments/Remarks），里面对驳回或 objection 提出争辩。",
          "amend_specification：A1 文件里有 SPEC（Specification），或 A... 修改书里有 'Amendments to the Specification'。",
          "amend_drawings：A1 文件里有 DRW（Drawings / Replacement Sheet）。",
          "interview：A1 文件里有请求面谈或面谈记录的文件（例如 M865 Letter Requesting Interview with Examiner）。REM 里'如有需要请电话联系'这种客套话不算。",
          "terminal_disclaimer：A1 文件里有 DIST.E.FILE / DISQ（Terminal Disclaimer）。",
          "declaration_affidavit：A1 文件里有 AF/D.130 / AF/D.131 / AF/D.132（Affidavit / Declaration）。",
          "",
          "自动标签对就填 ok；不对就填 fix，并写出正确的完整标签和依据。",
          "之后会按固定规则把这 9 类合成计划要求的 4 类：amendment = amend_claims/add_claims/amend_specification/amend_drawings；argument = argue；cancellation = cancel_claims；other = interview/terminal_disclaimer/declaration_affidavit。"]:
    g.append([t])
g.column_dimensions["A"].width = 160
wb.save(A / "gold" / "response_strategy_review_v2.xlsx"); print("ok", len(apps))
