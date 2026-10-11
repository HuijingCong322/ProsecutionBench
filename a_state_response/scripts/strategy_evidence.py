"""Write checkable evidence (file + line number + quoted line) for each strategy label into
gold/response_strategy_review_v2.xlsx column K, so the reviewer only verifies. Evidence is found by fixed rules:
amend_claims / add_claims: a line in the A1 claim listing labelled (Currently Amended) / (New).
cancel_claims: a claim labelled (Canceled) in the A1 listing that is listed and not canceled in the last listing before A1.
argue: REM lines stating traversal/argument (e.g. 'respectfully', 'traverse', 'submits').
other labels: the A1 document of the matching USPTO code (SPEC, DRW, M865/INTV, DIST.E.FILE/DISQ, AF/D.13x).
A label with no evidence is reported as 'NO EVIDENCE'; evidence for a label not in the automatic set is reported as 'MISSING LABEL'."""
import csv, glob, json, os, re
from pathlib import Path
from openpyxl import load_workbook
A = Path(__file__).resolve().parent.parent; R = A.parent; S = R / "data_v1_share"
LBL = re.compile(r"(?:Claims?\s+)?(\d{1,3})\s*\.?\s*(?:-\s*(\d{1,3})\s*\.?)?\s*\(\s*(currently\s+amended|previously\s+presented|original|new|canceled|cancelled|withdrawn)\s*\)", re.I)
NUM = re.compile(r"(?m)^\s*(\d{1,3})\s*\.\s+[A-Z]")
def listing(path):
    st, line = {}, {}
    txt = open(path).read()
    ms = list(LBL.finditer(txt))
    if not ms:   # claims as filed carry no status labels: numbered claims = original
        for m in NUM.finditer(txt):
            n = int(m.group(1)); st.setdefault(n, "original"); line.setdefault(n, txt[:m.start()].count("\n") + 1)
        return st, line, txt.splitlines()
    for m in ms:
        s = re.sub(r"\s+", " ", m.group(3).lower()).replace("cancelled", "canceled"); ln = txt[:m.start()].count("\n") + 1
        for n in range(int(m.group(1)), int(m.group(2) or m.group(1)) + 1): st.setdefault(n, s); line.setdefault(n, ln)
    return st, line, txt.splitlines()
def ranges(ns):
    ns = sorted(ns); out = []
    for n in ns:
        if out and n == out[-1][1] + 1: out[-1][1] = n
        else: out.append([n, n])
    return ", ".join(f"{a}" if a == b else f"{a}-{b}" for a, b in out)
p = A / "gold" / "response_strategy_review_v2.xlsx"; wb = load_workbook(p); ws = wb["strategy"]
for row in ws.iter_rows(min_row=2):
    app = str(row[0].value); e = json.loads((S / "examples" / f"{app}.json").read_text())
    cat = {r["doc_id"]: r for r in csv.DictReader(open(R / "data_v1" / "apps" / app / "catalog.csv"))}
    a1 = [(d, cat[d]["doc_code"]) for d in e["a1"]["doc_ids"] if d in cat]
    def f(code):
        for d, c in a1:
            if c == code:
                x = [y for y in glob.glob(str(S / "apps" / app / "text" / f"*{d}*.txt")) if "substantive" not in y]
                if x: return x[0]
    auto = set(e["labels"]["response_strategy"]); ev = {}
    clm = f("CLM")
    if clm:
        st, ln, lines = listing(clm); name = os.path.basename(clm)
        am = [n for n, s in st.items() if "currently amended" in s]; nw = [n for n, s in st.items() if s == "new"]
        if am: ev["amend_claims"] = f"D 列 CLM 第 {ln[min(am)]} 行: \"{lines[ln[min(am)]-1].strip()[:40]}\"；共 {len(am)} 条 Currently Amended（{ranges(am)}）"
        if nw: ev["add_claims"] = f"D 列 CLM 第 {ln[min(nw)]} 行: \"{lines[ln[min(nw)]-1].strip()[:40]}\"；New: {ranges(nw)}"
        prevs = sorted(x for x in glob.glob(str(S / "apps" / app / "text" / "*_CLM_*.txt")) if os.path.basename(x)[:10] < e["cutoff"])
        cn = [n for n, s in st.items() if "canceled" in s]
        if cn and prevs:
            pst, pln, _ = listing(prevs[-1])
            newc = [n for n in cn if n in pst and "canceled" not in pst[n]]
            gone = [n for n in cn if n not in pst]
            if newc: ev["cancel_claims"] = (f"D 列 CLM 第 {ln[min(newc)]} 行起标 Canceled: {ranges(newc)}；G 列旧清单里这些还在"
                                            f"（第 {pln[min(newc)]} 行: {min(newc)}. ({pst[min(newc)]})）")
            elif cn:
                o1 = [y for y in glob.glob(str(S / "apps" / app / "text" / f"*{e['o1']['doc_id']}*.txt")) if "substantive" not in y]
                ol = open(o1[0]).read().splitlines() if o1 else []
                k = next((i for i, l in enumerate(ol) if re.search(r"cancel", l, re.I)), None)
                was = [n for n in cn if n in pst and "canceled" in pst[n]]
                ev["cancel_claims_note"] = (f"A1 标 Canceled 的 {ranges(cn)} 在这次之前就已删除"
                                            + (f"（G 列旧清单里已标 Canceled: {ranges(was)}）" if was else "")
                                            + (f"（G 列旧清单里没有列出: {ranges(gone)}）" if gone else "")
                                            + (f"；O1 第 {k+1} 行: \"{ol[k].strip()[:60]}\"" if k is not None else ""))
    rem = f("REM")
    if rem:
        lines = open(rem).read().splitlines()
        k = next((i for i, l in enumerate(lines) if re.search(r"respectfully|traverse|submit|disagree", l, re.I)), None)
        if k is not None: ev["argue"] = f"E 列 REM 第 {k+1} 行: \"{lines[k].strip()[:80]}\""
    for code, lab in (("SPEC", "amend_specification"), ("DRW", "amend_drawings"), ("M865", "interview"), ("INTV.SUM.APP", "interview"),
                      ("DIST.E.FILE", "terminal_disclaimer"), ("DISQ", "terminal_disclaimer"), ("AF/D.130", "declaration_affidavit"),
                      ("AF/D.131", "declaration_affidavit"), ("AF/D.132", "declaration_affidavit")):
        if any(c == code for _, c in a1): ev.setdefault(lab, f"C 列有 {code} — {next(cat[d]['description'] for d, c in a1 if c == code)}")
    out = []
    for lab in sorted(auto | {k for k in ev if not k.endswith("_note")}):
        if lab in auto and lab in ev: out.append(f"{lab}: {ev[lab]}")
        elif lab in auto: out.append(f"{lab}: NO EVIDENCE — {ev.get(lab + '_note', '没找到对应证据')}")
        else: out.append(f"{lab}: MISSING LABEL — 有证据但自动标签里没有: {ev[lab]}")
    row[10].value = "\n".join(out)
wb.save(p)
for row in ws.iter_rows(min_row=2, values_only=True): print(row[0], "\n  " + row[10].replace("\n", "\n  "))
