"""Sheet B of gold_review_v3 -> gold_review_B.xlsx: drop exact-duplicate rows
(same app, item, answer, rule; decision copied at integration), and make
'statuses used' rows checkable by listing Office (OUTGOING) docs in (O1, cutoff]."""
import csv, json
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
R = Path(__file__).resolve().parents[2]; G = R / "A" / "gold"
src = load_workbook(G / "gold_review_v3.xlsx")["B 两边一致-确认"]
rows = list(src.iter_rows())
draft = {(x["app_id"], x["stage"][:2]): x["cutoff"] for x in map(json.loads, open(G / "state_gold_draft.jsonl"))}
def office_between(app, a, b):
    out = []
    for c in csv.DictReader(open(R / "data_v1" / "apps" / app / "catalog.csv")):
        if c["direction"] == "OUTGOING" and a < c["official_date"] <= b and c["doc_code"] not in ADMIN:
            out.append(f'{c["official_date"]} {c["doc_code"]}')
    nxt = sorted(c["official_date"] + " " + c["doc_code"] for c in csv.DictReader(open(R / "data_v1" / "apps" / app / "catalog.csv"))
                 if c["direction"] == "OUTGOING" and c["official_date"] > b and c["doc_code"] not in ADMIN)
    return out, (nxt[0] if nxt else "无")
ADMIN = {"N417", "NTC.PUB", "WFEE", "PGPUB", "ISSUE.NTF", "EGRANT.NTF"}  # receipts / publication notices, not examiner actions
wb = Workbook(); ws = wb.active; ws.title = "B 去重"
hdr = ["#", "原表#", "申请号", "阶段", "核对什么", "一致的答案（auto = model）", "规则答案", "证据链 / cutoff 检查",
       "原文", "PDF", "OA原句", "判断 ok/fix/状态对", "fix 时写正确值"]
ws.append(hdr)
seen, dropped, n = {}, [], 0
for r in rows[1:]:
    v = [c.value for c in r]
    if v[0] is None: continue
    key = (v[1], v[3], v[4], v[5])
    if key in seen: dropped.append((v[0], seen[key])); continue
    seen[key] = v[0]; n += 1
    ev = v[6]
    if "statuses used" in str(v[3]):
        o1 = draft[(v[1], "O1")]; cut = draft[(v[1], v[2])]
        docs, nxt = office_between(v[1], o1, cut)
        ev = (f"O1={o1}; cutoff={cut}; O1 之后到 cutoff 之间的审查员文件（不含电子收据/公开通知）: {', '.join(docs) or '无'}; "
              f"cutoff 之后第一份: {nxt}。无 → 全部 outstanding")
    ws.append([n, v[0], v[1], v[2], v[3], v[4], v[5], ev, v[7], v[8], v[9], None, None])
    for j, c in ((9, r[7]), (10, r[8])):
        if c.hyperlink: ws.cell(ws.max_row, j).hyperlink = c.hyperlink.target; ws.cell(ws.max_row, j).font = Font(color="0563C1", underline="single")
    if v[5] and v[5] != v[4]:
        for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor="FFF2CC")
# statuses rows from sheet A that the annotator could not check (no evidence shown there)
for app in ("14460009", "14838440"):
    o1 = draft[(app, "O1")]; cut = draft[(app, "A1")]
    docs, nxt = office_between(app, o1, cut); n += 1
    ws.append([n, "A表", app, "A1", "被驳回的 claims： statuses used", "outstanding", None,
               f"O1={o1}; cutoff={cut}; O1 之后到 cutoff 之间的审查员文件（不含电子收据/公开通知）: {', '.join(docs) or '无'}; "
               f"cutoff 之后第一份: {nxt}。无 → 全部 outstanding", None, None, None, None, None])
for col, w in zip("ABCDEFGHIJKLM", [4, 6, 11, 5, 26, 45, 45, 55, 6, 6, 50, 14, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
i = wb.create_sheet("说明")
for line in ["只核对：状态标签、驳回的 claim 号码。版本号不用管（由规则从状态标签算出）。",
             "判断列：ok = 答案对；状态对 = 状态对、版本没查；fix = 错，在后一列写正确值。",
             "黄色行 = 规则答案和 auto/model 不同，优先看。",
             "statuses used 行：看'证据链'列，如果 O1 之后到 cutoff 之间没有审查员文件，就填 ok。",
             f"已去掉 {len(dropped)} 行完全重复的行（整合时复制原行判定）：" + ", ".join(f"{a}←{b}" for a, b in dropped)]:
    i.append([line])
i.column_dimensions["A"].width = 120
wb.save(G / "gold_review_B.xlsx")
print("kept", n, "dropped", len(dropped))
