"""Integrate human review decisions into the final state gold (A).

Item universe : gold/state_gold_review_v2.csv (one row per snapshot x item; built by build_gold_review_v2.py
                from the automatic reference gold/state_gold_draft.jsonl and the frozen-prompt model output).
Decisions     : gold/state_gold_review_round1.xlsx (disagreeing structure rows, 8 apps)
                gold/gold_review_yellow_done.xlsx  (disagreeing amended-claim texts, 8 apps)
                gold/gold_review_v3_done.xlsx      (sheet A: all items of 2 new apps; sheet C: 20 sampled texts)
                gold/gold_review_B_done.xlsx       (agreeing items of 8 apps, de-duplicated; optional)
Decisions are matched to items BY VALUE (app, stage, item, auto string, model string), never by row position
alone, so a decision only applies to the exact answer the annotator saw.
Version protocol (state_annotation_guide.md): a decision 'status ok' (状态对) keeps the annotated statuses and
takes version_ids from the deterministic provenance rule (provenance.py).
Outputs       : gold/state_gold.jsonl, gold/state_gold_provenance.csv (one row per item: final value, decision,
                decision file/row, source), gold/state_gold_text_provenance.csv, printed coverage report."""
import csv, json, re, sys
from collections import defaultdict, Counter
from pathlib import Path
from openpyxl import load_workbook
A = Path(__file__).resolve().parent.parent; G = A / "gold"
sys.path.insert(0, str(A / "scripts")); import provenance as P  # noqa: E402
TWO = {"14460009", "14838440"}

def S(x): return "" if x is None else str(x).strip()
def ids(s):
    s = re.sub(r"\[[^\]]*\]", "", S(s))
    if s in ("", "(none)"): return set()
    return {n for a, b in re.findall(r"(\d+)(?:-(\d+))?", s) for n in range(int(a), int(b or a) + 1)}
def parse_summary(s):
    """'1-7, 11-19: previously_presented [as_filed]; 8-10: withdrawn [X]' -> {claim: (status, version)}"""
    out = {}
    for part in filter(None, (p.strip() for p in S(s).split(";"))):
        m = re.match(r"(.+?):\s*([a-z_]+)\s*\[([^\]]*)\]\s*$", part)
        if not m: raise ValueError(f"cannot parse summary part: {part!r}")
        for n in ids(m.group(1)): out[n] = (m.group(2), m.group(3))
    return out
def ntext(t): return re.sub(r"\s+", " ", S(t)).strip()

items = list(csv.DictReader(open(G / "state_gold_review_v2.csv")))
def kind(w):
    if w.startswith("entered_claims (") : return ("list", "entered_claims")
    if w.startswith("proposed_claims ("): return ("list", "proposed_claims")
    m = re.match(r"(entered|proposed)_claims claim (\d+) text", w)
    if m: return ("text", m.group(1) + "_claims", int(m.group(2)))
    m = re.match(r"rejection (\S+): claims", w)
    if m: return ("rej", m.group(1))
    if w == "rejection statuses used": return ("statuses",)
    raise ValueError(w)

dec = {}   # item index -> dict(decision, value, file, row, note)
def put(i, d, value, f, row, note=""):
    if i in dec and dec[i]["value"] != value:
        print(f"  CONFLICT item {i} {items[i]['app_id']} {items[i]['what']}: {dec[i]['file']} vs {f}")
    dec[i] = {"decision": d, "value": value, "file": f, "row": row, "note": note}

# ---- round1 (Chinese sheet; structure rows only; text rows there were redone in the yellow sheet)
KIND_ZH = [("申请人答复里提交的 claims", ("list", "proposed_claims")), ("审查员审查的那套 claims", ("list", "entered_claims"))]
r1 = list(load_workbook(G / "state_gold_review_round1.xlsx")["核对表"].iter_rows(values_only=True))
for j, x in enumerate(r1[1:], start=2):
    d = re.sub(r"[^a-z]", "", S(x[10]).lower())
    if not d: continue
    st = "O1_cutoff" if S(x[1]).startswith("O1") else "A1_cutoff"
    k = next((kk for zh, kk in KIND_ZH if zh in S(x[2])), None)
    if k is None:
        m = re.search(r"被以 ([0-9a-z]+)", S(x[2]))
        k = ("rej", m.group(1))
    val = {"auto": S(x[6]), "model": S(x[7]), "fix": S(x[11])}[d]
    hit = [i for i, it in enumerate(items) if it["app_id"] == S(x[0]) and it["stage"] == st and kind(it["what"]) == k
           and S(it["auto"]) == S(x[6]) and S(it["model"]) == S(x[7])]
    if len(hit) != 1:
        print(f"  round1 row {j} ({x[0]} {st} {k}) matched {len(hit)} items -> not applied"); continue
    put(hit[0], d, val, "state_gold_review_round1.xlsx", j)

# ---- sheet A of v3 (2 new apps): k-th TWO item in v2 order, verified by value
wa = list(load_workbook(G / "gold_review_v3_done.xlsx")["A 两件新数据"].iter_rows(values_only=True))
two = [i for i, it in enumerate(items) if it["app_id"] in TWO]
for j, x in enumerate(wa[1:], start=2):
    if x[0] is None: continue
    i = two[int(x[0]) - 1]; it = items[i]
    assert it["app_id"] == S(x[1]) and S(it["auto"]) == S(x[4]) and S(it["model"]) == S(x[5]), (j, x[:6])
    d = S(x[12]).lower()
    if not d: continue
    if "same and correct" in d: d = "auto"
    if d in ("auto", "model", "rule", "fix"):
        val = {"auto": S(x[4]), "model": S(x[5]), "rule": S(x[6]), "fix": S(x[13])}[d]
        put(i, d, val, "gold_review_v3_done.xlsx#A", j)
    else:
        print(f"  A row {j} decision {d!r} not usable ({it['app_id']} {it['stage']} {it['what']})")

# ---- sheet B (agreeing items, 8 apps), optional
bfile = G / "gold_review_B_done.xlsx"
v3b = list(load_workbook(G / "gold_review_v3.xlsx")["B 两边一致-确认"].iter_rows(values_only=True))
agree = [i for i, it in enumerate(items) if it["app_id"] not in TWO and it["agree"] == "yes"]
v3b_item = {int(x[0]): agree[int(x[0]) - 1] for x in v3b[1:] if x[0] is not None}
for n, i in v3b_item.items():
    assert S(items[i]["auto"]) == S(v3b[n][4]), n
if bfile.exists():
    wb_ = list(load_workbook(bfile)["B 去重"].iter_rows(values_only=True))
    kept = {}
    for j, x in enumerate(wb_[1:], start=2):
        if x[0] is None: continue
        d = S(x[11]).lower(); fixv = S(x[12])
        if d.startswith("ok"): d = "ok"
        elif d not in ("fix", "状态对"): d = ""          # '?' or blank = no decision
        if S(x[1]) == "A表":   # statuses rows of the 2 new apps
            tg = [i for i, it in enumerate(items) if it["app_id"] == S(x[2]) and it["stage"].startswith(S(x[3]))
                  and it["what"] == "rejection statuses used"]
        else:
            tg = [v3b_item[int(x[1])]]
        for i in tg:
            if not d: continue
            v = fixv if d == "fix" else S(items[i]["auto"])
            put(i, "status_ok" if d == "状态对" else ("fix" if d == "fix" else "ok"), v, "gold_review_B_done.xlsx", j)
        if S(x[1]) != "A表": kept[(S(x[2]), S(x[4]), S(x[5]), S(x[6]))] = (tg[0], j)
    for n, i in v3b_item.items():      # dropped duplicates inherit
        x = v3b[n]; key = (S(x[1]), S(x[3]), S(x[4]), S(x[5]))
        if i not in dec and key in kept and kept[key][0] in dec:
            src = dec[kept[key][0]]
            put(i, src["decision"], src["value"], src["file"], src["row"], f"copied from identical item (B row {kept[key][1]})")
else:
    print("  (gold_review_B_done.xlsx not found: agreeing items of 8 apps stay unreviewed)")

CAT = {}
def catalog(app):
    if app not in CAT: CAT[app] = list(csv.DictReader(open(A.parent / "data_v1" / "apps" / app / "catalog.csv")))
    return CAT[app]
ADMIN = {"N417", "NTC.PUB", "WFEE", "PGPUB", "ISSUE.NTF", "EGRANT.NTF"}
draft_cut = {(x["app_id"], x["stage"]): x["cutoff"] for x in map(json.loads, open(G / "state_gold_draft.jsonl"))}
for i, it in enumerate(items):
    app = it["app_id"]
    # A1 'statuses used': outstanding iff no examiner document between O1 and the A1 cutoff (USPTO file wrapper catalog)
    if it["what"] == "rejection statuses used" and it["stage"] == "A1_cutoff" and i not in dec:
        o1, cut = draft_cut[(app, "O1_cutoff")], draft_cut[(app, "A1_cutoff")]
        docs = [c["official_date"] + " " + c["doc_code"] for c in catalog(app)
                if c["direction"] == "OUTGOING" and o1 < c["official_date"] <= cut and c["doc_code"] not in ADMIN]
        if not docs and S(it["auto"]) == "outstanding":
            put(i, "catalog_check", "outstanding", "data_v1/apps/%s/catalog.csv" % app, "-",
                f"no examiner document in ({o1}, {cut}]; rejections cannot have been withdrawn or overcome")
        else:
            print(f"  statuses {app} A1: examiner docs before cutoff {docs} -> needs human decision")
    # A1 proposed list empty: the A1 reply contains no claim listing (CLM) on the cutoff date
    if it["what"].startswith("proposed_claims (") and it["stage"] == "A1_cutoff" and i not in dec and not S(it["auto"]) and not S(it["model"]):
        cut = draft_cut[(app, "A1_cutoff")]
        if not [c for c in catalog(app) if c["official_date"] == cut and c["doc_code"] == "CLM"]:
            put(i, "doc_check", "", "data_v1/apps/%s/catalog.csv" % app, "-", f"A1 reply of {cut} has no claim listing (CLM)")

# O1 'statuses used': at the O1 cutoff every rejection was issued that day -> outstanding by definition
for i, it in enumerate(items):
    if it["what"] == "rejection statuses used" and it["stage"] == "O1_cutoff" and i not in dec and S(it["auto"]) == "outstanding":
        put(i, "definitional", "outstanding", "-", "-", "O1 snapshot: rejections issued at the cutoff are outstanding")
    if it["what"].startswith("proposed_claims (") and it["stage"] == "O1_cutoff" and i not in dec and not S(it["auto"]) and not S(it["model"]):
        put(i, "definitional", "", "-", "-", "O1 snapshot: no applicant response yet, proposed list is empty")
# A sheet: identical O1/A1 rejection rows of the same app were judged separately; nothing to copy.

# ---- text decisions (value-keyed): yellow, C, A-text rows
txt = {}   # (app, claim, normalized auto text) -> (decision, value, file, row)
yl = list(load_workbook(G / "gold_review_yellow_done.xlsx")["黄色行（claim文本）"].iter_rows(values_only=True))
for j, x in enumerate(yl[1:], start=2):
    if x[0] is None or not S(x[6]): continue
    d = re.sub(r"[^a-z]", "", S(x[6]).lower()); v = {"auto": S(x[8]), "model": S(x[9]), "fix": S(x[7])}[d]
    txt[(S(x[1]), int(x[2]), ntext(x[8]))] = (d, v, "gold_review_yellow_done.xlsx", j)
wc = list(load_workbook(G / "gold_review_v3_done.xlsx")["C 文本抽查20"].iter_rows(values_only=True))
for j, x in enumerate(wc[1:], start=2):
    if x[0] is None or not S(x[8]): continue
    d = re.sub(r"[^a-z]", "", S(x[8]).lower()); v = S(x[9]) if d == "fix" else S(x[4])
    txt[(S(x[1]), int(x[3]), ntext(x[4]))] = (d, v, "gold_review_v3_done.xlsx#C", j)
for i, it in enumerate(items):
    k = kind(it["what"])
    if k[0] == "text" and i in dec:
        txt[(it["app_id"], k[2], ntext(it["auto"]))] = (dec[i]["decision"], dec[i]["value"], dec[i]["file"], dec[i]["row"])
# items of kind text are carried by the text table; mark covered text items
for i, it in enumerate(items):
    k = kind(it["what"])
    if k[0] == "text" and i not in dec:
        t = txt.get((it["app_id"], k[2], ntext(it["auto"])))
        if t: dec[i] = {"decision": t[0], "value": t[1], "file": t[2], "row": t[3], "note": "text table"}

# ---- build snapshots
draft = {(x["app_id"], x["stage"]): x for x in map(json.loads, open(G / "state_gold_draft.jsonl"))}
pred = {(x["app_id"], x["stage"]): x["state"] for x in map(json.loads, open(A / "outputs" / "state_predictions.jsonl"))
        if x.get("group") == "gold" and x.get("parse_ok")}
rules = {app: P.stage_rule(app, "") for app in {k[0] for k in draft}}
prov, tprov, gold, issues = [], [], [], []
for (app, stage), d in sorted(draft.items()):
    its = [(i, it) for i, it in enumerate(items) if it["app_id"] == app and it["stage"] == stage]
    rec = {"sample_id": d["sample_id"], "app_id": app, "stage": stage, "cutoff": d["cutoff"],
           "entered_claims": [], "proposed_claims": [], "rejections": []}
    for i, it in its:
        k = kind(it["what"]); dd = dec.get(i)
        final = dd["value"] if dd else S(it["auto"])
        src = dd["decision"] if dd else "auto_unreviewed"
        rule_val = ""
        if k[0] == "list":
            rl = "entered" if k[1] == "entered_claims" else "proposed"
            rule_val = "" if (rl == "proposed" and stage.startswith("O1")) else rules[app][rl][0]
            claims = parse_summary(final)
            if src == "status_ok":
                rv = parse_summary(rule_val)
                claims = {n: (s, rv.get(n, (None, v))[1]) for n, (s, v) in claims.items()}
            rv = parse_summary(rule_val) if rule_val else {}
            vdiff = sorted(n for n in claims if n in rv and rv[n][1] != claims[n][1])
            if vdiff: issues.append(f"{app} {stage} {k[1]}: version differs from rule for claims {vdiff}")
            dl = {c["claim_id"]: c for c in d[k[1]]}
            pl = {int(c["claim_id"]): c for c in (pred.get((app, stage)) or {}).get(k[1], []) if str(c.get("claim_id")).isdigit()}
            for n in sorted(claims):
                s, v = claims[n]
                base = dl.get(n, {}).get("text") if n in dl else None; tsrc = "auto_unreviewed"
                if base is None:
                    if s == "canceled": base, tsrc = "", "canceled"
                    elif n in pl: base, tsrc = pl[n].get("text", ""), "model (claim absent from auto reference)"
                    else: base, tsrc = "", "MISSING"; issues.append(f"{app} {stage} {k[1]} claim {n}: no text")
                t = txt.get((app, n, ntext(base)))
                if t and s != "canceled": text, tsrc = t[1], f"{t[0]} ({t[2]} row {t[3]})"
                else: text = base
                c2 = re.sub(r"^\s*IN\s*T\s*H\s*E\s+CLAIMS\s*:?\s*(\d{1,3}\s*)?\d{1,3}\s*\.\s*\([A-Za-z ]+\)\s*", "", text or "")
                if c2 != (text or ""): tsrc += " + listing header/label removed"
                rec[k[1]].append({"claim_id": n, "status": s, "version_id": v, "text": c2})
                tprov.append({"app_id": app, "stage": stage, "list": k[1], "claim_id": n, "status": s,
                              "version_id": v, "text_source": tsrc})
        elif k[0] == "rej":
            ev = next((r.get("evidence") for r in d["rejections"] if r["ground"] == k[1]), [])
            for n in sorted(ids(final)):
                rec["rejections"].append({"claim_id": n, "ground": k[1], "status": None, "evidence": ev})
        prov.append({"app_id": app, "stage": stage, "item": it["what"], "auto": it["auto"][:300], "model": it["model"][:300],
                     "rule": rule_val, "final": final[:600], "decision": src, "decision_file": dd["file"] if dd else "",
                     "decision_row": dd["row"] if dd else "", "note": dd["note"] if dd else "",
                     "evidence_doc": it.get("doc_id", "")})
    stat = next((dec[i]["value"] if i in dec else S(it["auto"]) for i, it in its if it["what"] == "rejection statuses used"), "outstanding")
    if "," in stat: issues.append(f"{app} {stage}: several statuses used ({stat}); per-claim status taken from auto")
    for r in rec["rejections"]:
        r["status"] = stat if "," not in stat else next((x["status"] for x in d["rejections"]
                                                         if x["claim_id"] == r["claim_id"] and x["ground"] == r["ground"]), "outstanding")
    gold.append(rec)

with open(G / "state_gold.jsonl", "w") as f:
    for r in gold: f.write(json.dumps(r, ensure_ascii=False) + "\n")
for name, rows in (("state_gold_provenance.csv", prov), ("state_gold_text_provenance.csv", tprov)):
    with open(G / name, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
c = Counter(p["decision"] for p in prov)
print("snapshots", len(gold), "| items", len(prov), "| decisions:", dict(c))
print("text records by source:", dict(Counter(t["text_source"].split(" (")[0] for t in tprov)))
for s in issues: print("  ISSUE", s)
