"""Gold review sheet v3 (reproducible provenance). Sheets:
 A: the two apps whose inputs changed in v1.2 (14460009, 14838440): every item, with candidate documents,
    the rule-derived answer (A/scripts/provenance.py) and links;
 B: items where the automatic reference and the model agree (other 8 apps): quick confirmation;
 C: 20 randomly sampled agreeing claim texts (seed 20261009) for a text spot check against the PDF.
Writes A/gold/gold_review_v3.xlsx and A/gold/pdf_needed.txt (documents whose PDF the reviewer needs)."""
import csv, json, random, re, difflib, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
A = Path(__file__).resolve().parents[1]; R = A.parent
sys.path.insert(0, str(A / "scripts"))
import provenance as P  # noqa: E402

TWO = {"14460009", "14838440"}
rows = list(csv.DictReader(open(A / "gold" / "state_gold_review_v2.csv")))
date = {}
def ddate(app, did):
    if (app, did) not in date:
        for c in csv.DictReader(open(R / "data_v1" / "apps" / app / "catalog.csv")):
            date[(app, c["doc_id"])] = c["official_date"]
    return date.get((app, did), "")
def textlink(app, did):
    if not did: return ""
    f = [p for p in (R / "data_v1" / "apps" / app / "text").glob(f"*{did}*.txt") if "substantive" not in p.name]
    return f"file:///Users/fionay/examinerbench/{f[0].relative_to(R)}" if f else ""
need = set()
def pdflink(app, did):
    need.add((app, did)); return f"file:///Users/fionay/examinerbench/A/gold/pdf/{app}_{ddate(app, did)}_{did}.pdf"
rule = {}
def R_(app):
    if app not in rule: rule[app] = P.stage_rule(app, "")
    return rule[app]
def diff(a, b):
    A_, B_ = re.findall(r"\S+", a), re.findall(r"\S+", b); out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, A_, B_, autojunk=False).get_opcodes():
        if op == "equal": continue
        if i2 > i1: out.append("【auto独有: " + " ".join(A_[i1:i2]) + "】")
        if j2 > j1: out.append("【model独有: " + " ".join(B_[j1:j2]) + "】")
    return "\n".join(out)[:1500]
def norm(x): return re.sub(r"\s+", "", x or "")

wb = Workbook()
intro = wb.active; intro.title = "说明"
for t in [
 "【来源规则（论文附录同款，脚本 A/scripts/provenance.py 可复现）】",
 "审查员审查的 claims（entered）= O1 之前提交的所有 claim 文件（CLM），按日期依次叠加；同一天先原始、后修改。",
 "申请人提出的 claims（proposed）= 上面的结果再叠加 A1 答复里的 claim 文件。",
 "每条 claim：状态 = 最近一份列出它的文件里的标注；版本号 = 最近一次把它 Currently Amended / New / 首次 Canceled 的文件编号；从没改过 = as_filed。",
 "某份文件没有列出之前存在的 claim 时，规则会在'规则提示'里标出，需要人工依据其他文件（PGPUB、后续 listing）裁定。",
 "",
 "【怎么做】",
 "A 页（2 件新数据）：和之前一样，判断填 auto / model / rule / fix（rule = 采用'规则答案'列）。fix 时写正确值。证据链列出了所有候选文件。",
 "B 页（两边一致）：看'一致的答案'是否与原文相符；与'规则答案'不同的行标黄，要重点看。判断填 ok 或 fix。",
 "C 页（文本抽查 20 条）：打开 PDF 找到这条 claim，看文本是否正确（删除线去掉、下划线保留）。判断填 ok 或 fix。",
 "PDF 需先下载：cd ~/examinerbench && source ~/.zshrc && python3 A/scripts/fetch_review_pdfs.py"]:
    intro.append([t])
intro.column_dimensions["A"].width = 150

yellow = PatternFill("solid", fgColor="FFF4CE"); red = PatternFill("solid", fgColor="FDE2E1"); head = PatternFill("solid", fgColor="E8EEF7")
def sheet(title, hdr, widths):
    ws = wb.create_sheet(title); ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True); c.fill = head
    for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "D2"; return ws
def link(cell, url, label):
    if url: cell.value = label; cell.hyperlink = url; cell.font = Font(color="0563C1", underline="single")
def finish(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

# ---- sheet A
wa = sheet("A 两件新数据", ["#", "申请号", "阶段", "核对什么", "auto", "model", "规则答案", "证据链（所有候选 claim 文件）", "规则提示",
                          "原文", "PDF", "两边不同的词 / OA原句", "判断 auto/model/rule/fix", "fix 时写正确值"],
           [4, 10, 6, 26, 40, 40, 40, 55, 30, 8, 8, 50, 12, 40])
k = 0
for r in rows:
    if r["app_id"] not in TWO: continue
    k += 1; app = r["app_id"]; st = "O1" if r["stage"].startswith("O1") else "A1"; w = r["what"]; src = None
    lst = "entered" if w.startswith("entered") else "proposed" if w.startswith("proposed") else None
    rr = R_(app)
    if lst and "text" not in w:
        what = "审查员审查的 claims（状态 [版本号]）" if lst == "entered" else "申请人答复里的 claims（状态 [版本号]）"
        ra, chain, notes = rr[lst][0], rr[lst][1], "\n".join(rr[lst][2])
        if lst == "proposed" and st == "O1": ra, chain, notes = "", "O1 阶段：还没有申请人答复，应为空", ""
        extra = ""
    elif lst:
        n = int(re.search(r"claim (\d+)", w).group(1)); what = f"claim {n} 的文本"
        src = rr["text_source"][lst].get(n); ra = src[2][:1500] if src else ""
        chain = f"规则：此 claim 的文本来自 {src[1]} CLM {src[0]}" if src else ""; notes = ""; extra = diff(r["auto"], r["model"])
    else:
        what = r["what"].replace("rejection", "被驳回的 claims：").replace(": claims", ""); ra = ""; chain = ""; notes = ""; extra = r.get("evidence", "")
    ws = wa; ws.append([k, app, st, what, r["auto"], r["model"], ra, chain, notes, "", "", extra, "", ""])
    did = r["doc_id"] or (src[0] if lst and "text" in w and src else "")
    link(ws.cell(ws.max_row, 10), textlink(app, did), "文本")
    if did: link(ws.cell(ws.max_row, 11), pdflink(app, did), "PDF")
    if r["agree"] == "NO":
        for c in ws[ws.max_row][:9]: c.fill = red
finish(wa)

# ---- sheet B
wbk = sheet("B 两边一致-确认", ["#", "申请号", "阶段", "核对什么", "一致的答案（auto = model）", "规则答案", "证据链", "原文", "PDF", "OA原句", "判断 ok/fix", "fix 时写正确值"],
            [4, 10, 6, 26, 50, 50, 50, 8, 8, 50, 10, 40])
k = 0
for r in rows:
    if r["app_id"] in TWO or r["agree"] != "yes": continue
    k += 1; app = r["app_id"]; st = "O1" if r["stage"].startswith("O1") else "A1"; w = r["what"]
    lst = "entered" if w.startswith("entered") else "proposed" if w.startswith("proposed") else None
    if lst:
        rr = R_(app); ra, chain = rr[lst][0], rr[lst][1] + ("\n" + "\n".join(rr[lst][2]) if rr[lst][2] else "")
        if lst == "proposed" and st == "O1":
            ra, chain = "", "O1 阶段：还没有申请人答复，应为空"
        what = "审查员审查的 claims（状态 [版本号]）" if lst == "entered" else "申请人答复里的 claims（状态 [版本号]）"
    else:
        ra = chain = ""; what = w.replace("rejection", "被驳回的 claims：").replace(": claims", "")
    wbk.append([k, app, st, what, r["auto"], ra, chain, "", "", r.get("evidence", ""), "", ""])
    link(wbk.cell(wbk.max_row, 8), textlink(app, r["doc_id"]), "文本")
    if r["doc_id"]: link(wbk.cell(wbk.max_row, 9), pdflink(app, r["doc_id"]), "PDF")
    def ids(x): return {int(n) for a, b in re.findall(r"(\d+)(?:-(\d+))?", re.sub(r"\[[^\]]*\]", "", x)) for n in range(int(a), int(b or a) + 1)}
    if lst and ra and ids(ra) != ids(r["auto"]):
        wbk.cell(wbk.max_row, 7).value = (chain + "\n" if chain else "") + "【注意】规则未能从原文解析出全部编号（多为 OCR 质量差），以原文为准"
    elif lst and norm(ra) != norm(r["auto"]):
        for c in wbk[wbk.max_row][:7]: c.fill = yellow
finish(wbk)

# ---- sheet C: 20 sampled agreeing claim texts
draft = {(x["app_id"], x["stage"]): x for x in map(json.loads, open(A / "gold" / "state_gold_draft.jsonl"))}
pred = {(x["app_id"], x["stage"]): x for x in map(json.loads, open(A / "outputs" / "state_predictions.jsonl")) if x.get("group") == "gold"}
pool = []
for (app, stage), d in draft.items():
    if app in TWO or (app, stage) not in pred or not pred[(app, stage)].get("state"): continue
    for lst in ("entered_claims", "proposed_claims"):
        pm = {c["claim_id"]: c for c in pred[(app, stage)]["state"].get(lst, [])}
        for c in d.get("state", d)[lst]:
            q = pm.get(c["claim_id"])
            if not q or c["status"] == "canceled" or not c.get("text"): continue
            if difflib.SequenceMatcher(None, norm(c["text"]).lower(), norm(q.get("text")).lower(), autojunk=False).ratio() >= 0.9:
                pool.append((app, stage, lst, c["claim_id"], c["text"]))
seen = set(); uniq = []
for x in sorted(pool):
    if (x[0], x[3], x[4]) in seen: continue
    seen.add((x[0], x[3], x[4])); uniq.append(x)
sample = random.Random(20261009).sample(uniq, min(20, len(uniq)))
wc = sheet("C 文本抽查20", ["#", "申请号", "阶段", "claim", "文本（auto 与 model 一致）", "来源文件（规则）", "原文", "PDF", "判断 ok/fix", "fix 时写正确值"],
           [4, 10, 6, 6, 80, 30, 8, 8, 10, 40])
for i, (app, stage, lst, n, txt) in enumerate(sorted(sample), 1):
    src = R_(app)["text_source"]["entered" if lst == "entered_claims" else "proposed"].get(n)
    did = src[0] if src else ""
    wc.append([i, app, "O1" if stage.startswith("O1") else "A1", n, txt, f"{src[1]} CLM {did}" if src else "", "", "", "", ""])
    link(wc.cell(wc.max_row, 7), textlink(app, did), "文本")
    if did: link(wc.cell(wc.max_row, 8), pdflink(app, did), "PDF")
finish(wc)
wb.save(A / "gold" / "gold_review_v3.xlsx")
(A / "gold" / "pdf_needed.txt").write_text("\n".join(f"{a}\t{d}" for a, d in sorted(need)))
print("A rows", wa.max_row - 1, "| B rows", wbk.max_row - 1, "| C rows", wc.max_row - 1, "| pool", len(uniq), "| PDFs needed", len(need))
